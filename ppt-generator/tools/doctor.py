#!/usr/bin/env python3
"""环境能力探测。

设计原则：探测，不依赖。任何一个可选组件缺失都不应导致失败，
只降级并明确告知使用者「少了什么 / 影响什么 / 装什么能补回来」。

核心教训（2026-07 实测）：命令存在 != 能用。
本机 soffice 存在，却因看不到中文字体而静默产出「没有中文」的 PDF。
所以渲染能力必须用带中文的测试页做端到端验证，数中文区域像素，而不是 which soffice。

用法：
    python3 doctor.py            # 人类可读报告
    python3 doctor.py --json     # 机器可读，供生成流程判断降级
"""

import glob
import json
import os
import shutil
import subprocess
import sys
import tempfile

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ── 渲染器候选路径：按顺序探测，找到即用 ───────────────────────
# 注意：不硬编码任何特定 runtime 的路径为「依赖」。
# codex 缓存那条属于「顺手捡」——有就用，没有完全无所谓。
SOFFICE_CANDIDATES = [
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
    "/opt/homebrew/bin/soffice",
    os.path.expanduser(
        "~/.cache/codex-runtimes/*/dependencies/native/libreoffice-headless/"
        "libreoffice/*.app/Contents/MacOS/soffice"
    ),
    os.path.expanduser("~/.cache/codex-runtimes/*/dependencies/bin/override/soffice"),
]

# ── 中文字体候选：用于 PIL 度量。skill 自带优先，保证跨机一致 ──
FONT_CANDIDATES = [
    os.path.join(SKILL_DIR, "assets", "fonts", "NotoSansSC-Regular.otf"),
    os.path.join(SKILL_DIR, "assets", "fonts", "NotoSansSC-Regular.ttf"),
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    os.path.expanduser("~/Library/Fonts/NotoSansSC-Regular.otf"),
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "C:/Windows/Fonts/msyh.ttc",
]


def _resolve(candidates):
    """支持 glob 的路径解析。返回第一个可执行/存在的。"""
    for c in candidates:
        if "*" in c:
            for hit in sorted(glob.glob(c)):
                if os.access(hit, os.X_OK):
                    return hit
        elif os.path.sep in c:
            if os.access(c, os.X_OK):
                return c
        else:
            found = shutil.which(c)
            if found:
                return found
    return None


def check_pptx():
    try:
        import pptx  # noqa: F401
        return True, getattr(pptx, "__version__", "?")
    except ImportError:
        return False, None


def check_pillow():
    try:
        import PIL
        return True, getattr(PIL, "__version__", "?")
    except ImportError:
        return False, None


def check_font():
    """找一个 PIL 能真正加载的中文字体（存在 != 能加载，PingFang.ttc 就加载不了）。"""
    try:
        from PIL import ImageFont
    except ImportError:
        return None
    for path in FONT_CANDIDATES:
        if not os.path.exists(path):
            continue
        try:
            f = ImageFont.truetype(path, 20)
            if f.getlength("客户增长") > 0:   # 必须真能量出中文宽度
                return path
        except Exception:
            continue
    return None


def check_renderer():
    """端到端验证渲染链路：pptx -> pdf -> png，且中文必须真的画出来。

    返回 (soffice路径, pdf2png工具, 中文是否OK, 诊断信息)
    """
    soffice = _resolve(SOFFICE_CANDIDATES)
    if not soffice:
        return None, None, False, "未找到 LibreOffice"

    pdf2png = shutil.which("pdftoppm") or shutil.which("magick") or shutil.which("convert")
    if not pdf2png:
        return soffice, None, False, "找到 LibreOffice，但缺 PDF 转图工具"

    ok_pptx, _ = check_pptx()
    if not ok_pptx:
        return soffice, pdf2png, False, "缺 python-pptx，无法生成测试页"

    # ── 端到端冒烟测试：画一页中文，渲染后数中文区域的墨迹像素 ──
    from pptx import Presentation
    from pptx.util import Inches, Pt

    tmp = tempfile.mkdtemp(prefix="pptdoctor_")
    try:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(2))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = "中文渲染冒烟测试客户增长效果验证"
        r.font.size = Pt(60)
        src = os.path.join(tmp, "smoke.pptx")
        prs.save(src)

        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
            capture_output=True, timeout=180,
        )
        pdf = os.path.join(tmp, "smoke.pdf")
        if not os.path.exists(pdf):
            return soffice, pdf2png, False, "LibreOffice 未能产出 PDF"

        if "pdftoppm" in pdf2png:
            subprocess.run([pdf2png, "-png", "-r", "60", pdf, os.path.join(tmp, "p")],
                           capture_output=True, timeout=120)
        else:
            subprocess.run([pdf2png, "-density", "60", pdf, os.path.join(tmp, "p-1.png")],
                           capture_output=True, timeout=120)

        pngs = glob.glob(os.path.join(tmp, "p*.png"))
        if not pngs:
            return soffice, pdf2png, False, "PDF 转 PNG 失败（可能缺 ghostscript，建议装 poppler）"

        from PIL import Image
        im = Image.open(pngs[0]).convert("L")
        # 文字区域（上半部）的非白像素数。中文若被丢弃，这里会接近 0。
        crop = im.crop((0, 0, im.width, im.height // 2))
        hist = crop.histogram()
        ink = sum(hist[:128])
        if ink < 500:
            return soffice, pdf2png, False, (
                f"渲染器看不到中文字体（墨迹仅 {ink}px）——会静默产出无中文的预览。"
                f"把中文字体放进 LibreOffice 的 Contents/Resources/fonts/truetype 可修复"
            )
        return soffice, pdf2png, True, f"端到端可信（中文墨迹 {ink}px）"
    except subprocess.TimeoutExpired:
        return soffice, pdf2png, False, "渲染超时"
    except Exception as e:
        return soffice, pdf2png, False, f"渲染冒烟测试异常：{e}"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def check_icons():
    d = os.path.join(SKILL_DIR, "assets", "icons")
    return len(glob.glob(os.path.join(d, "*.png"))) if os.path.isdir(d) else 0


def collect():
    ok_pptx, v_pptx = check_pptx()
    ok_pil, v_pil = check_pillow()
    font = check_font()
    soffice, pdf2png, render_ok, render_msg = check_renderer()
    return {
        "can_generate": ok_pptx,
        "pptx_version": v_pptx,
        "pillow": v_pil,
        "font_for_metrics": font,
        "soffice": soffice,
        "pdf2png": pdf2png,
        "render_qa_ok": render_ok,
        "render_msg": render_msg,
        "bundled_icons": check_icons(),
    }


def report(c):
    L = []
    A = L.append
    A("═" * 62)
    A("  ppt-generator · 环境能力报告")
    A("═" * 62)

    A("")
    A("【必需】生成 PPT")
    if c["can_generate"]:
        A(f"  ✅ python-pptx {c['pptx_version']}")
    else:
        A("  ❌ 缺 python-pptx —— 无法生成，请先： pip install python-pptx")
        A("")
        A("  ⛔ 缺少必需依赖，生成流程无法继续。")
        A("═" * 62)
        return "\n".join(L), "blocked"

    A(f"  ✅ 内置图标 {c['bundled_icons']} 个（无需联网、无需 librsvg）")

    A("")
    A("【建议】文字度量 · 防溢出")
    if c["font_for_metrics"]:
        A(f"  ✅ Pillow {c['pillow']} + {os.path.basename(c['font_for_metrics'])}")
        A("     → 按真实字体度量排版，溢出可精确拦截")
    else:
        A("  ⚠️  无可用中文字体（或缺 Pillow）")
        A("     → 降级为宽度模型估算（中文按 1.0em），精度下降，长文本可能溢出")
        A("     → 恢复：把中文字体放到 skill 的 assets/fonts/")

    A("")
    A("【建议】渲染自审 · 落差最大的一档")
    if c["render_qa_ok"]:
        A(f"  ✅ {os.path.basename(c['soffice'])} + {os.path.basename(c['pdf2png'])}")
        A(f"     → {c['render_msg']}")
        A("     → 生成后会导出每页 PNG 与 montage，逐页复核换行/重叠/拥挤/配色")
    else:
        A(f"  ⚠️  渲染自审不可用：{c['render_msg']}")
        A("     → 影响：无法看见成品，只能盲拼坐标。版式正确但无人复核，")
        A("       换行、重叠、拥挤、观感问题不保证被发现。")
        A("     → 恢复：brew install --cask libreoffice && brew install poppler")
        A("             （Linux: apt install libreoffice poppler-utils）")

    A("")
    A("─" * 62)
    if c["render_qa_ok"] and c["font_for_metrics"]:
        A("  结论：能力完整 —— 生成 + 度量 + 渲染自审全部可用。")
        level = "full"
    elif c["font_for_metrics"] or c["render_qa_ok"]:
        A("  结论：能力部分降级 —— 可生成，复核能力不完整（见上）。")
        level = "partial"
    else:
        A("  结论：仅可生成 —— 无度量、无自审。")
        A("  ⚠️  产出的是「设计系统正确但未经复核」的 PPT，建议自行检查排版。")
        level = "minimal"
    A("═" * 62)
    return "\n".join(L), level


if __name__ == "__main__":
    caps = collect()
    if "--json" in sys.argv:
        text, level = report(caps)
        caps["level"] = level
        print(json.dumps(caps, ensure_ascii=False, indent=2))
    else:
        text, level = report(caps)
        print(text)
    sys.exit(0 if caps["can_generate"] else 1)
